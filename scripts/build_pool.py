import sys,re; sys.stdout.reconfigure(encoding='utf-8')
import warnings; warnings.filterwarnings('ignore')
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
WORK=os.environ.get('JOBHUNT_HOME','./workspace')
B=WORK+'/'
d=pd.read_excel(B+'数据/杭州岗位评分_v2_20261023.xlsx',sheet_name='岗位评分')
S=lambda v:'' if pd.isna(v) else str(v)

# —— 已核验的结论（来自原文核对/爱企查）——
ai_bad=set()  # 可选：把已人工核验为不合适的岗位链接放进来
bad_names=[l.strip() for l in open(B+'blacklist.txt',encoding='utf-8') if l.strip()] if os.path.exists(B+'blacklist.txt') else []  # 你核验后确认不投的公司关键词，一行一个
nontech=r'漫剧|短剧|视频|美工|导演|编剧|剪辑|主播|直播|运营|文案|设计师|制作师|生成师|BD|商务|销售|客服|培训|讲师|审核|标注|招聘|猎头|产品助理|售前|地推|招生|美术|特效|动效|原画|建模|策划|技术美术'
x=d[d['排除标记'].isna()&d['外包嫌疑'].isna()].copy()
x=x[~x['岗位'].str.contains(nontech,regex=True,na=False)]
x=x[~x['岗位'].str.contains(r'2[789]届|2[789]秋招|2[789]春招|日语|派遣|外派|wai派|OD',regex=True,na=False)]
x=x[~x['学历'].fillna('').str.contains('硕士|博士')]
x=x[x['岗位'].str.contains(r'开发|工程师|研发|程序|全栈|前端|后端|实施|Unity|U3D|客户端|Agent|智能体|测试',regex=True,na=False)]
x=x[~x['经验'].fillna('').str.contains(r'3-5年|5-10年|10年以上|3年以上|5年以上')]
x=x[~x['岗位'].str.contains(r'高级|资深|专家|架构|经理|负责人|主管|组长',regex=True,na=False)]
x=x[~x['链接'].isin(ai_bad)]
x=x[~x['公司'].str.contains('|'.join(bad_names),regex=True,na=False)]
x=x[~((x['赛道']=='实习')&(x['毕业生可投']=='存疑(仅限在校)'))]
# 基本真实性：太小/作坊式名称先不进池
x=x[~x['规模'].fillna('').str.contains('0-20人|^20人以下|1-49人',regex=True)]
x=x[~x['公司'].str.contains('工作室|商行|个体|经营部|劳务|人力|猎头',regex=True,na=False)]
x=x[~x['公司'].str.startswith('某')]
def cat(t):
    if re.search(r'Unity|U3D|游戏',t,re.I): return 'Unity/游戏'
    if re.search(r'前端|Vue|Web|H5',t,re.I): return '前端'
    if re.search(r'实施|交付',t): return '实施'
    if re.search(r'AI|大模型|Agent|智能体',t,re.I): return 'AI'
    return '其他开发/测试'
x['方向']=x['岗位'].map(cat)
ver={'Unity/游戏':'Unity游戏版','前端':'前端版','实施':'通用版','AI':'AI实施版(需先做完Agent项目)','其他开发/测试':'通用版'}
x['简历版本']=x['方向'].map(ver)
def tier(r):
    up=re.findall(r'\d+',S(r['规模']))
    up=max(map(int,up)) if up else 0
    if up>=1000: return '冲(大厂/大公司，靠运气/内推)'
    return '海投'
x['层']=x.apply(tier,axis=1)
# 精投层：已爱企查核验的真实公司
good=[l.strip() for l in open(B+'whitelist.txt',encoding='utf-8') if l.strip()] if os.path.exists(B+'whitelist.txt') else []  # 你核验后确认真实的公司关键词
x.loc[x['公司'].str.contains('|'.join(good),na=False),'层']='精投(已核验)'
# 一个公司在同一方向只留得分最高的一个岗位（同平台只能投一次）
x=x.sort_values('总分',ascending=False).drop_duplicates(['公司','方向'])
order={'精投(已核验)':0,'海投':1,'冲(大厂/大公司，靠运气/内推)':2}
x['_o']=x['层'].map(order)
x=x.sort_values(['_o','总分'],ascending=[True,False]).reset_index(drop=True)
x['批次']=(x.index//25+1)
# AI岗位还没做项目 → 放到最后批次
x.loc[x['方向']=='AI','批次']=x['批次'].max()+1
cols=['批次','层','方向','简历版本','公司','岗位','薪资','规模','经验','学历','平台','机器评分','链接']
x['机器评分']=x['总分']
for c in ['投递日期','是否已问HR(接受毕业生)','结果(无回复/约面/拒/offer)','备注']: x[c]=''
out=x[cols+['投递日期','是否已问HR(接受毕业生)','结果(无回复/约面/拒/offer)','备注']]
f=B+'海投池_20261023.xlsx'
with pd.ExcelWriter(f,engine='openpyxl') as w:
    out.to_excel(w,sheet_name='海投池',index=False)
    out[out['批次']==1].to_excel(w,sheet_name='第1批(25家)',index=False)
    pd.DataFrame({'说明':[
     '这是"可以投"的池子，不是"一定靠谱"的名单。已去掉：外包/派遣、非技术岗、27届校招、日语、硕博、高级岗、要3年以上、公司太小或作坊式名称、已核验为不对口/有风险的。',
     '层：精投=已用爱企查核验的真实公司（定制简历和话术）；海投=批量投；冲=大公司，当抽奖。',
     '批次：每批25个，建议每天投一批。AI方向岗位放在最后（要先做完Agent项目再投）。',
     '同一公司同一方向只留一个岗位（同平台同公司短期只能投一次）。',
     '投之前：红字补完；每个"实习"岗先问HR"接受2026届已毕业的同学吗"；Boss的岗位还没抓全，后面补。',
     '用法：他每次投完在"投递日期/结果"里记一下，这张表就是投递跟踪表。',
     '风险提醒：海投池没有逐家做工商核验。面试邀约来了，再查那一家（爱企查看参保人数、经营范围）。']}).to_excel(w,sheet_name='说明',index=False)
wb=load_workbook(f)
for ws in wb:
    ws.freeze_panes='A2'
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E79')
    for i,col in enumerate(ws.columns,1):
        ml=max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:60])
        ws.column_dimensions[get_column_letter(i)].width=min(max(8,ml*1.4),50)
        for c in col: c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.auto_filter.ref=ws.dimensions
wb.save(f)
print('池子',len(out),'个岗位',out['公司'].nunique(),'家公司')
print(out.groupby('层').agg(岗位=('岗位','size')).to_string())
print(out.groupby('方向').agg(岗位=('岗位','size')).to_string())
print('批次数',out['批次'].max())
print(out[out['批次']==1][['层','方向','公司','岗位','薪资','规模']].head(12).to_string(index=False))
