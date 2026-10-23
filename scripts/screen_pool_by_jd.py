# 海投池二次筛选：读岗位原文，把"特别差"的踢掉，并写明原因
import sys, re, json, glob
sys.stdout.reconfigure(encoding='utf-8')
import warnings; warnings.filterwarnings('ignore')
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import os
WORK = os.environ.get('JOBHUNT_HOME', './workspace')
B = WORK + '/'
pool = pd.read_excel(B + '海投池_20261023.xlsx', sheet_name='海投池')

def clip(t):
    t = re.sub(r'\s+', ' ', t)
    return re.split(r'猎聘温馨提示|公司简介|相似职位|公司信息|工作地点|关于我们|职位发布者', t)[0]

KILL_REGEX = [  # (正则, 原因) —— 命中即踢
    (r'(?<![\d.])([2-9]|1\d)\s*年(?:以上|及以上|或以上|工作经验|相关(?:工作)?经验|开发经验|前端经验|经验)', '要2年以上经验'),
    (r'丰富.{0,8}经验|经验丰富|资深', '要丰富经验/资深'),
    (r'硕士(?:及以上|学历|以上|优先)?(?!.{0,6}本科)', None),   # 单独处理
    (r'博士', '要博士'),
    (r'日语|N[12]\b', '要日语'),
    (r'2[789]届|2[789]秋招|2[789]春招', '27届及以后校招'),
    (r'外派|外包|驻场|wai派|派遣|OD岗|德科|项目外包', '外派/外包'),
    (r'培训贷|分期|押金|保证金|服装费|培训费|收费|先培训|包分配|无责底薪|底薪.{0,6}提成|日结|刷单|代理|加盟|拉人|兼职', '套路/兼职/收费类'),
    (r'电话(?:销售|邀约)|拓客|地推|客户开发|业务拓展|销售(?:额|业绩)|业绩(?:提成|指标)', '销售伪装成技术岗'),
]

def screen(text, job):
    t = clip(text)
    reasons = []
    for rx, why in KILL_REGEX:
        m = re.search(rx, job + ' ' + t)
        if not m: continue
        if why is None:
            if re.search(r'硕士', t) and not re.search(r'本科', t): reasons.append('要硕士')
        else:
            reasons.append(why)
    if len(re.sub(r'\W', '', t)) < 35: reasons.append('岗位描述过短/空')
    return sorted(set(reasons))

def biz_check(biz, comp):
    """智联页面里的工商信息：虚报规模嫌疑 / 新公司提示"""
    flags = []
    cap = re.search(r'注册资本\s*([\d,.]+)\s*万', biz)
    size = re.findall(r'\d+', re.search(r'(\d+-\d+人|\d+人以上|\d+人以下)', comp).group(0)) if re.search(r'(\d+-\d+人|\d+人以上|\d+人以下)', comp) else []
    up = max(map(int, size)) if size else 0
    found = re.search(r'成立时间\s*(\d{4})-(\d{2})', biz)
    kill = None
    if cap and float(cap.group(1).replace(',', '')) <= 10 and up >= 100: kill = f'疑似虚报规模：注册资本{cap.group(1)}万，页面写{up}人'
    if found and (int(found.group(1)), int(found.group(2))) >= (2026, 4): flags.append('成立不足半年')
    return kill, flags

res = {}
for f in glob.glob(B + '数据/raw/research_pool_*.json'):
    for o in json.load(open(f, encoding='utf-8')).get('out', []):
        if 'text' in o: res[o['u']] = o

rows = []
for r in pool.itertuples():
    url = r.链接
    verdict, why, flags = '未核验(Boss/读不到原文)', [], []
    if url in res:
        o = res[url]
        why = screen(o['text'], o['co'] and r.岗位)
        kill, fl = biz_check(o.get('biz', ''), o.get('comp', ''))
        flags += fl
        if kill: why.append(kill)
        verdict = '踢掉' if why else '保留(原文已核)'
    rows.append((verdict, '；'.join(why), '；'.join(flags)))
pool['核验'] = [x[0] for x in rows]; pool['踢掉原因'] = [x[1] for x in rows]; pool['提示'] = [x[2] for x in rows]

kept = pool[pool['核验'] != '踢掉'].copy()
kicked = pool[pool['核验'] == '踢掉'].copy()
kept['批次'] = (kept.reset_index(drop=True).index // 25 + 1).values
kept.loc[kept['方向'] == 'AI', '批次'] = kept['批次'].max() + 1
kept = kept.sort_values(['批次']).reset_index(drop=True)

cols = ['批次', '层', '核验', '提示', '方向', '简历版本', '公司', '岗位', '薪资', '规模', '经验', '学历', '平台', '机器评分', '链接',
        '投递日期', '是否已问HR(接受毕业生)', '结果(无回复/约面/拒/offer)', '备注']
out = B + '海投池_v2_已踢差评_20261023.xlsx'
with pd.ExcelWriter(out, engine='openpyxl') as w:
    kept[cols].to_excel(w, sheet_name='海投池(保留)', index=False)
    kept[kept['批次'] == 1][cols].to_excel(w, sheet_name='第1批', index=False)
    kicked[['层', '方向', '公司', '岗位', '薪资', '规模', '踢掉原因', '链接']].to_excel(w, sheet_name='已踢掉(含原因)', index=False)
    pd.DataFrame({'说明': [
        '在海投池基础上，读了岗位原文，把特别差的踢掉：要2年以上经验/资深、硕博、日语、27届校招、外派外包、套路/兼职/收费类、销售伪装成技术岗、描述过短，以及智联页面上"注册资本≤10万却写100人以上"的疑似虚报规模。',
        '"保留(原文已核)"=读了原文没发现硬伤；"未核验"=Boss直聘的岗位，当前读不到原文（风控），投之前他自己点开看一眼岗位要求。',
        '"提示"列：成立不足半年等，不踢但要留意。',
        '被踢掉的都列在"已踢掉(含原因)"，万一规则误伤可以捞回来。']}).to_excel(w, sheet_name='说明', index=False)
wb = load_workbook(out)
for ws in wb:
    ws.freeze_panes = 'A2'
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79')
    for i, col in enumerate(ws.columns, 1):
        ml = max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:60])
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, ml * 1.4), 50)
        for c in col: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.auto_filter.ref = ws.dimensions
wb.save(out)
print('原池', len(pool), '→ 保留', len(kept), '，踢掉', len(kicked), '，其中未核验(Boss)', (kept['核验'].str.startswith('未核验')).sum())
print('踢掉原因分布:'); print(kicked['踢掉原因'].str.split('；').explode().value_counts().to_string())
print('\n保留中各层:', kept['层'].value_counts().to_dict())
print('保留中各方向:', kept['方向'].value_counts().to_dict())
print('\n被踢掉的示例:'); print(kicked[['公司', '岗位', '踢掉原因']].head(12).to_string(index=False))
